"""Dependency-light legacy SSE MCP server for remote document generation."""

from __future__ import annotations

import json
import mimetypes
import os
import queue
import re
import secrets
import threading
import uuid
from datetime import datetime
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, quote, urlparse

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


SERVER_NAME = "weknora-document-editor-sse"
PROTOCOL_VERSION = "2025-03-26"
OUTPUT_DIR = Path(os.getenv("OUTPUT_DIR", "/data/generated")).resolve()
PUBLIC_BASE_URL = os.getenv("PUBLIC_BASE_URL", "").rstrip("/")
API_KEY = os.getenv("MCP_API_KEY", "")
PORT = int(os.getenv("PORT", "7860"))
SESSIONS: dict[str, queue.Queue[dict[str, Any]]] = {}
SESSIONS_LOCK = threading.Lock()


def schema(properties: dict[str, Any], required: list[str]) -> dict[str, Any]:
    return {"type": "object", "properties": properties, "required": required, "additionalProperties": False}


TABLE_SCHEMA = {"type": "array", "items": {"type": "object", "properties": {"headers": {"type": "array", "items": {"type": "string"}}, "rows": {"type": "array", "items": {"type": "array", "items": {}}}}, "required": ["headers", "rows"], "additionalProperties": False}}
TOOLS = [
    {"name": "create_word", "description": "Create a downloadable .docx document.", "inputSchema": schema({"filename": {"type": "string"}, "title": {"type": "string"}, "paragraphs": {"type": "array", "items": {"type": "string"}}, "tables": TABLE_SCHEMA}, ["filename"])},
    {"name": "create_excel", "description": "Create a downloadable .xlsx workbook.", "inputSchema": schema({"filename": {"type": "string"}, "sheets": {"type": "array", "items": {"type": "object", "properties": {"name": {"type": "string"}, "headers": {"type": "array", "items": {"type": "string"}}, "rows": {"type": "array", "items": {"type": "array", "items": {}}}}, "required": ["name", "headers", "rows"], "additionalProperties": False}}}, ["filename", "sheets"])},
    {"name": "create_pdf", "description": "Create a downloadable Unicode-capable .pdf report.", "inputSchema": schema({"filename": {"type": "string"}, "title": {"type": "string"}, "paragraphs": {"type": "array", "items": {"type": "string"}}, "tables": TABLE_SCHEMA}, ["filename"])},
]


def text(value: Any) -> str:
    return "" if value is None else str(value)


def path_for(filename: str, suffix: str) -> tuple[str, str, Path]:
    stem = re.sub(r"[^A-Za-z0-9._ -]", "_", Path(text(filename)).stem).strip(" .") or "document"
    token = secrets.token_urlsafe(18)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    public_name = f"{stem}{suffix}"
    return token, public_name, OUTPUT_DIR / f"{token}--{public_name}"


def file_result(path: Path, token: str, public_name: str, base_url: str, mime_type: str) -> dict[str, Any]:
    url = f"{base_url}/files/{token}/{quote(public_name)}"
    return {"content": [{"type": "text", "text": f"Created {public_name}. Attach the download URL to the conversation."}, {"type": "resource_link", "uri": url, "name": public_name, "description": f"Generated at {datetime.now().isoformat(timespec='seconds')}", "mimeType": mime_type}], "structuredContent": {"url": url, "filename": public_name, "mimeType": mime_type}}


def create_word(arguments: dict[str, Any], base_url: str) -> dict[str, Any]:
    token, public_name, path = path_for(arguments["filename"], ".docx")
    document = Document()
    document.sections[0].top_margin = document.sections[0].bottom_margin = Inches(0.75)
    if arguments.get("title"):
        document.add_heading(text(arguments["title"]), 0)
    for paragraph in arguments.get("paragraphs", []):
        document.add_paragraph(text(paragraph))
    for spec in arguments.get("tables", []):
        table = document.add_table(rows=1, cols=max(1, len(spec["headers"])))
        table.style = "Table Grid"
        for index, header in enumerate(spec["headers"]):
            table.rows[0].cells[index].text = text(header)
        for row in spec["rows"]:
            cells = table.add_row().cells
            for index, value in enumerate(row[:len(cells)]):
                cells[index].text = text(value)
    document.save(path)
    return file_result(path, token, public_name, base_url, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


def create_excel(arguments: dict[str, Any], base_url: str) -> dict[str, Any]:
    token, public_name, path = path_for(arguments["filename"], ".xlsx")
    workbook = Workbook()
    workbook.remove(workbook.active)
    for spec in arguments["sheets"]:
        sheet = workbook.create_sheet(re.sub(r"[\\/*?:\[\]]", "_", text(spec["name"]))[:31] or "Sheet")
        sheet.append([text(header) for header in spec["headers"]])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row in spec["rows"]:
            sheet.append(list(row))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(50, max(12, max(len(text(cell.value)) for cell in column) + 2))
    workbook.save(path)
    return file_result(path, token, public_name, base_url, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def create_pdf(arguments: dict[str, Any], base_url: str) -> dict[str, Any]:
    token, public_name, path = path_for(arguments["filename"], ".pdf")
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontName="STSong-Light", fontSize=10, leading=15)
    heading = ParagraphStyle("Title", parent=styles["Title"], fontName="STSong-Light", fontSize=18, leading=24)
    story: list[Any] = []
    if arguments.get("title"):
        story.extend([Paragraph(text(arguments["title"]), heading), Spacer(1, 6 * mm)])
    for paragraph in arguments.get("paragraphs", []):
        story.extend([Paragraph(text(paragraph).replace("\n", "<br/>"), body), Spacer(1, 3 * mm)])
    for spec in arguments.get("tables", []):
        rows = [[text(value) for value in spec["headers"]]] + [[text(value) for value in row] for row in spec["rows"]]
        table = Table(rows, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), "STSong-Light"), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9CA3AF")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.extend([table, Spacer(1, 5 * mm)])
    SimpleDocTemplate(str(path), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm).build(story or [Paragraph(" ", body)])
    return file_result(path, token, public_name, base_url, "application/pdf")


HANDLERS = {"create_word": create_word, "create_excel": create_excel, "create_pdf": create_pdf}


class Handler(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"

    def log_message(self, format: str, *args: Any) -> None:
        print(format % args, flush=True)

    def base_url(self) -> str:
        if PUBLIC_BASE_URL:
            return PUBLIC_BASE_URL
        proto = self.headers.get("X-Forwarded-Proto", "http").split(",")[0]
        return f"{proto}://{self.headers.get('Host', f'localhost:{PORT}') }"

    def allowed(self) -> bool:
        return not API_KEY or secrets.compare_digest(self.headers.get("Authorization", "").removeprefix("Bearer "), API_KEY)

    def json_response(self, status: int, payload: dict[str, Any]) -> None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(body)

    def do_OPTIONS(self) -> None:
        self.send_response(HTTPStatus.NO_CONTENT)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Headers", "Authorization, Content-Type")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.end_headers()

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/health":
            self.json_response(HTTPStatus.OK, {"status": "ok", "server": SERVER_NAME})
            return
        if parsed.path.startswith("/files/"):
            self.serve_file(parsed.path)
            return
        if parsed.path != "/sse" or not self.allowed():
            self.json_response(HTTPStatus.UNAUTHORIZED if parsed.path == "/sse" else HTTPStatus.NOT_FOUND, {"error": "Unauthorized" if parsed.path == "/sse" else "Not found"})
            return
        session_id = uuid.uuid4().hex
        messages: queue.Queue[dict[str, Any]] = queue.Queue()
        with SESSIONS_LOCK:
            SESSIONS[session_id] = messages
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "text/event-stream")
        self.send_header("Cache-Control", "no-cache")
        self.send_header("Connection", "keep-alive")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(f"event: endpoint\ndata: {self.base_url()}/messages?sessionId={session_id}\n\n".encode())
        self.wfile.flush()
        try:
            while True:
                try:
                    message = messages.get(timeout=15)
                    payload = json.dumps(message, ensure_ascii=False)
                    self.wfile.write(f"event: message\ndata: {payload}\n\n".encode())
                except queue.Empty:
                    self.wfile.write(b": keepalive\n\n")
                self.wfile.flush()
        except (BrokenPipeError, ConnectionResetError):
            pass
        finally:
            with SESSIONS_LOCK:
                SESSIONS.pop(session_id, None)

    def do_POST(self) -> None:
        parsed = urlparse(self.path)
        session_id = parse_qs(parsed.query).get("sessionId", [""])[0]
        if parsed.path != "/messages" or not self.allowed():
            self.json_response(HTTPStatus.UNAUTHORIZED, {"error": "Unauthorized"})
            return
        with SESSIONS_LOCK:
            messages = SESSIONS.get(session_id)
        if messages is None:
            self.json_response(HTTPStatus.NOT_FOUND, {"error": "Unknown SSE session"})
            return
        try:
            request = json.loads(self.rfile.read(int(self.headers.get("Content-Length", "0"))))
        except (ValueError, json.JSONDecodeError):
            self.json_response(HTTPStatus.BAD_REQUEST, {"error": "Invalid JSON"})
            return
        result = self.handle_mcp(request)
        if result is not None:
            messages.put(result)
        self.send_response(HTTPStatus.ACCEPTED)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Content-Length", "0")
        self.end_headers()

    def serve_file(self, request_path: str) -> None:
        parts = request_path.split("/")
        if len(parts) != 4 or not re.fullmatch(r"[A-Za-z0-9_-]{20,}", parts[2]):
            self.json_response(HTTPStatus.NOT_FOUND, {"error": "Not found"})
            return
        filename = Path(parts[3]).name
        candidate = (OUTPUT_DIR / f"{parts[2]}--{filename}").resolve()
        if candidate.parent != OUTPUT_DIR or not candidate.is_file():
            self.json_response(HTTPStatus.NOT_FOUND, {"error": "Not found"})
            return
        content = candidate.read_bytes()
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", mimetypes.guess_type(candidate.name)[0] or "application/octet-stream")
        self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    def handle_mcp(self, request: dict[str, Any]) -> dict[str, Any] | None:
        request_id, method = request.get("id"), request.get("method")
        if method == "notifications/initialized":
            return None
        if method == "initialize":
            result: dict[str, Any] = {"protocolVersion": PROTOCOL_VERSION, "capabilities": {"tools": {"listChanged": False}}, "serverInfo": {"name": SERVER_NAME, "version": "1.0.0"}}
        elif method == "tools/list":
            result = {"tools": TOOLS}
        elif method == "tools/call":
            params = request.get("params", {})
            handler = HANDLERS.get(params.get("name"))
            if handler is None:
                return {"jsonrpc": "2.0", "id": request_id, "error": {"code": -32602, "message": "Unknown tool"}}
            try:
                result = handler(params.get("arguments", {}), self.base_url())
            except (KeyError, TypeError, ValueError) as exc:
                result = {"content": [{"type": "text", "text": f"Invalid tool input: {exc}"}], "isError": True}
            except Exception as exc:
                result = {"content": [{"type": "text", "text": f"Document generation failed: {exc}"}], "isError": True}
        else:
            return {"jsonrpc": "2.0", "id": request_id, "error": {"code": -32601, "message": "Method not found"}}
        return {"jsonrpc": "2.0", "id": request_id, "result": result}


if __name__ == "__main__":
    print(f"{SERVER_NAME} listening on 0.0.0.0:{PORT}", flush=True)
    ThreadingHTTPServer(("0.0.0.0", PORT), Handler).serve_forever()
